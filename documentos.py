"""Geração de documentos oficiais de secretaria escolar (memorando, ofício, etc.)."""

import docx
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONTE = "Arial"
TAM_NORMAL = 12
TAM_TITULO = 16

LARGURAS_COLUNAS_CM = [3.8, 7.4, 3.8, 3.5]  # soma = 18.5cm = largura util da pagina A4 (margens estreitas)
ALTURA_MINIMA_CABECALHO_CM = 1.7
ALTURA_MINIMA_DEPARA_CM = 1.3
ALTURA_MINIMA_ASSUNTO_CM = 1.8
ALTURA_MINIMA_CORPO_CM = 11.0  # garante que a pagina fique cheia mesmo com pouco texto
                                # (deixa espaco pro rodape crescer quando tem imagem de assinatura)
CHARS_POR_LINHA_CORPO = 85  # estimativa de quantos caracteres cabem numa linha do corpo
ALTURA_LINHA_CM = 0.53  # altura aproximada de uma linha de texto a 12pt


def _estimar_linhas(texto, chars_por_linha=CHARS_POR_LINHA_CORPO):
    """Estimativa grosseira de quantas linhas um paragrafo vai ocupar quando
    quebrado automaticamente pelo Word - usada so pra calcular quanto espaco
    em branco falta antes do fecho, nao precisa ser exata."""
    texto = texto.strip()
    if not texto:
        return 0
    return max(1, -(-len(texto) // chars_por_linha))


def _definir_margens_celulas(tabela, cima_cm=0.25, baixo_cm=0.25, esquerda_cm=0.2, direita_cm=0.2):
    """Da mais respiro interno as celulas (por padrao o python-docx nao coloca quase nada)."""
    tblPr = tabela._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for nome, valor in (("top", cima_cm), ("bottom", baixo_cm), ("left", esquerda_cm), ("right", direita_cm)):
        el = OxmlElement(f"w:{nome}")
        el.set(qn("w:w"), str(int(Cm(valor).twips)))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def _altura_minima(linha_tabela, altura_cm):
    linha_tabela.height = Cm(altura_cm)
    linha_tabela.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST


def _config_secao(document):
    secao = document.sections[0]
    secao.page_width = Cm(21)
    secao.page_height = Cm(29.7)
    secao.left_margin = Cm(1.5)
    secao.right_margin = Cm(1.0)
    secao.top_margin = Cm(1.8)
    secao.bottom_margin = Cm(1.0)


def _remover_bordas_tabela(tabela):
    tbl = tabela._tbl
    tblPr = tbl.tblPr
    bordas = OxmlElement("w:tblBorders")
    for nome in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{nome}")
        el.set(qn("w:val"), "nil")
        bordas.append(el)
    tblPr.append(bordas)


def _definir_largura_colunas(tabela, larguras_cm):
    tabela.autofit = False
    tabela.allow_autofit = False
    for linha in tabela.rows:
        for celula, largura in zip(linha.cells, larguras_cm):
            celula.width = Cm(largura)
    for idx, largura in enumerate(larguras_cm):
        tabela.columns[idx].width = Cm(largura)


def _set_fonte_padrao(document):
    estilo = document.styles["Normal"]
    estilo.font.name = FONTE
    estilo.font.size = Pt(TAM_NORMAL)
    estilo.paragraph_format.space_before = Pt(0)
    estilo.paragraph_format.space_after = Pt(0)


def _paragrafo(celula_ou_doc, texto="", negrito=False, tamanho=TAM_NORMAL,
               alinhamento=None, indice=0):
    if hasattr(celula_ou_doc, "paragraphs") and hasattr(celula_ou_doc, "add_paragraph"):
        if indice == 0 and celula_ou_doc.paragraphs and not celula_ou_doc.paragraphs[0].runs:
            p = celula_ou_doc.paragraphs[0]
        else:
            p = celula_ou_doc.add_paragraph()
    else:
        p = celula_ou_doc.add_paragraph()
    if alinhamento is not None:
        p.alignment = alinhamento
    if texto:
        run = p.add_run(texto)
        run.font.name = FONTE
        run.font.size = Pt(tamanho)
        run.bold = negrito
    return p


def _impedir_quebra_de_linha(linha_tabela):
    """Evita que uma linha de tabela seja partida entre duas páginas."""
    trPr = linha_tabela._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    trPr.append(cant_split)


def _imagem_centralizada(destino, caminho_imagem, largura_cm):
    if not caminho_imagem:
        return None
    p = destino.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(caminho_imagem, width=Cm(largura_cm))
    return p


def _celula_vazia(celula):
    """Centraliza verticalmente uma celula recem-criada (ja vem sem
    conteudo, nao precisa limpar nada - so ajustar o alinhamento).

    Nota: NAO fazer `paragrafo.text = ""` aqui - no python-docx 1.2.0 isso
    cria um run vazio de verdade, o que faz _paragrafo() achar que a celula
    ja tem conteudo e pular pra um paragrafo novo, deixando uma linha em
    branco indesejada antes do primeiro texto real."""
    celula.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    return celula


def _linhas_cabecalho(dados_escola):
    """Linhas de texto do cabecalho que aparece abaixo do brasao/logo - nome
    da escola e secretaria/orgao superior, iguais no Memorando, no Oficio e
    na Declaracao."""
    linhas = []
    nome_escola = (dados_escola.get("nome_escola") or "").strip()
    if nome_escola:
        linhas.append(nome_escola)
    secretaria = (dados_escola.get("secretaria") or "").strip()
    if secretaria:
        linhas.append(secretaria)
    return linhas


def _linha_contato(dados_escola):
    """Uma linha combinando endereco/telefone/email da escola, pra completar
    o cabecalho do Memorando e do Oficio - a Declaracao ja mostra esses dados
    numa tabela propria, entao nao usa esta linha (evita duplicar)."""
    partes = []
    endereco = (dados_escola.get("endereco") or "").strip()
    telefone = (dados_escola.get("telefone") or "").strip()
    email = (dados_escola.get("email") or "").strip()
    if endereco:
        partes.append(endereco)
    if telefone:
        partes.append(f"Tel: {telefone}")
    if email:
        partes.append(email)
    return " - ".join(partes)


def gerar_memorando(dados_escola, dados_memo, caminho_saida):
    """Gera um Memorando (.docx) a partir dos dados da escola e do memorando.

    Todo o documento fica em uma unica tabela com bordas visiveis, separando
    cada bloco (cabecalho, DE/PARA, ASSUNTO, corpo, rodape) em quadros, igual
    ao modelo de memorando usado nas secretarias escolares. Recebido por e a
    segunda Data ficam em branco no documento, para preencher a mao. O resto
    ja vem preenchido pelo que o usuario digitou/cadastrou.

    dados_escola: dict com nome_escola, secretaria, diretor_nome,
        diretor_cargo, diretor_portaria, logo_path, assinatura_path.
    dados_memo: dict com numero, ano, para, assunto, saudacao, corpo (texto
        com quebras de linha), fecho, assinado_por (nome de quem assina -
        so aparece na caixa ASSINATURA quando a escola nao tem uma imagem de
        assinatura cadastrada), protocolo (texto livre, ex: "SIGED Nº 12345"
        - se vazio, a linha de protocolo nem aparece no cabecalho), data.
    """
    document = docx.Document()
    _config_secao(document)
    _set_fonte_padrao(document)

    protocolo = (dados_memo.get("protocolo") or "").strip()
    usa_protocolo = bool(protocolo)
    linhas_cabecalho = 2 if usa_protocolo else 1
    total_linhas = linhas_cabecalho + 1 + 1 + 1 + 1  # cabecalho + DE/PARA + ASSUNTO + corpo + rodape

    tabela = document.add_table(rows=total_linhas, cols=4)
    tabela.style = "Table Grid"
    _definir_largura_colunas(tabela, LARGURAS_COLUNAS_CM)
    _definir_margens_celulas(tabela)

    for linha in tabela.rows[:linhas_cabecalho]:
        _altura_minima(linha, ALTURA_MINIMA_CABECALHO_CM)

    # --- Cabecalho: logo + titulo (+ protocolo, se a escola usar) ---
    if usa_protocolo:
        cel_logo = _celula_vazia(tabela.cell(0, 0).merge(tabela.cell(1, 0)))
        cel_titulo = _celula_vazia(tabela.cell(0, 1).merge(tabela.cell(0, 3)))
        cel_protocolo = _celula_vazia(tabela.cell(1, 1).merge(tabela.cell(1, 3)))
    else:
        cel_logo = _celula_vazia(tabela.cell(0, 0))
        cel_titulo = _celula_vazia(tabela.cell(0, 1).merge(tabela.cell(0, 3)))

    if dados_escola.get("logo_path"):
        _imagem_centralizada(cel_logo, dados_escola["logo_path"], 2.8)

    for i, linha in enumerate(_linhas_cabecalho(dados_escola)):
        _paragrafo(cel_titulo, linha, negrito=(i == 0), tamanho=13 if i == 0 else TAM_NORMAL)
    linha_contato = _linha_contato(dados_escola)
    if linha_contato:
        _paragrafo(cel_titulo, linha_contato, tamanho=TAM_NORMAL - 1)

    secretaria = (dados_escola.get("secretaria") or "").strip()
    numero = (dados_memo.get("numero") or "").strip()
    ano = (dados_memo.get("ano") or "").strip()
    titulo_memo = f"MEMORANDO Nº {numero}/{ano}" if numero else "MEMORANDO Nº"
    if secretaria:
        titulo_memo += f" - {secretaria}"
    _paragrafo(cel_titulo, titulo_memo, negrito=True, tamanho=TAM_TITULO)

    if usa_protocolo:
        _paragrafo(cel_protocolo, protocolo, negrito=True)

    proxima_linha = linhas_cabecalho

    # --- DE / PARA ---
    _altura_minima(tabela.rows[proxima_linha], ALTURA_MINIMA_DEPARA_CM)
    cel_de = _celula_vazia(tabela.cell(proxima_linha, 0).merge(tabela.cell(proxima_linha, 1)))
    cel_para = _celula_vazia(tabela.cell(proxima_linha, 2).merge(tabela.cell(proxima_linha, 3)))
    _paragrafo(cel_de, f"DE: {dados_escola.get('nome_escola', '')}", negrito=True)
    para = (dados_memo.get("para") or "").strip()
    _paragrafo(cel_para, f"PARA: {para}" if para else "PARA:", negrito=True)
    proxima_linha += 1

    # --- ASSUNTO ---
    _altura_minima(tabela.rows[proxima_linha], ALTURA_MINIMA_ASSUNTO_CM)
    cel_assunto = _celula_vazia(tabela.cell(proxima_linha, 0).merge(tabela.cell(proxima_linha, 3)))
    cel_assunto.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    p_assunto = cel_assunto.paragraphs[0]
    r1 = p_assunto.add_run("ASSUNTO: ")
    r1.bold = True
    r1.font.name = FONTE
    r1.font.size = Pt(TAM_NORMAL)
    assunto = (dados_memo.get("assunto") or "").strip()
    if assunto:
        r2 = p_assunto.add_run(assunto.upper())
        r2.font.name = FONTE
        r2.font.size = Pt(TAM_NORMAL)
    proxima_linha += 1

    # --- Corpo do memorando ---
    _altura_minima(tabela.rows[proxima_linha], ALTURA_MINIMA_CORPO_CM)
    cel_corpo = _celula_vazia(tabela.cell(proxima_linha, 0).merge(tabela.cell(proxima_linha, 3)))
    cel_corpo.vertical_alignment = WD_ALIGN_VERTICAL.TOP

    p_saud = cel_corpo.paragraphs[0]
    p_saud.paragraph_format.first_line_indent = Cm(1.25)
    run_saud = p_saud.add_run(dados_memo.get("saudacao") or "Prezado(a) Senhor(a),")
    run_saud.font.name = FONTE
    run_saud.font.size = Pt(TAM_NORMAL)
    linhas_usadas = 1

    cel_corpo.add_paragraph()
    linhas_usadas += 1

    corpo = dados_memo.get("corpo") or ""
    paragrafos_corpo = [linha.strip() for linha in corpo.split("\n") if linha.strip()]
    for linha in paragrafos_corpo:
        p = cel_corpo.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(1.25)
        run = p.add_run(linha)
        run.font.name = FONTE
        run.font.size = Pt(TAM_NORMAL)
        linhas_usadas += _estimar_linhas(linha)

    cel_corpo.add_paragraph()
    linhas_usadas += 1

    # Empurra o fecho pra mais perto do rodape (como no modelo original),
    # preenchendo com linhas em branco o espaco que sobrar dentro da altura
    # minima do corpo - sem isso, textos curtos deixam o fecho "colado" logo
    # depois do ultimo paragrafo, bem no topo da caixa.
    altura_util_corpo_cm = ALTURA_MINIMA_CORPO_CM - 0.5  # desconta margens da celula
    linhas_disponiveis = altura_util_corpo_cm / ALTURA_LINHA_CM
    linhas_preenchimento = int(linhas_disponiveis - linhas_usadas - 1)
    linhas_preenchimento = max(0, min(linhas_preenchimento, 20))
    for _ in range(linhas_preenchimento):
        cel_corpo.add_paragraph()

    p_fecho = cel_corpo.add_paragraph()
    p_fecho.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_fecho = p_fecho.add_run(dados_memo.get("fecho") or "Atenciosamente,")
    run_fecho.font.name = FONTE
    run_fecho.font.size = Pt(TAM_NORMAL)
    proxima_linha += 1

    # --- Rodape: DATA | ASSINATURA | RECEBIDO POR | DATA ---
    linha_rodape = tabela.rows[proxima_linha]
    _impedir_quebra_de_linha(linha_rodape)
    cel_data, cel_assinatura, cel_recebido, cel_data2 = linha_rodape.cells
    for c in (cel_data, cel_assinatura, cel_recebido, cel_data2):
        _celula_vazia(c)

    # RECEBIDO POR e a segunda DATA ficam em branco de proposito: sao
    # preenchidos a mao na hora de entregar/receber o memorando.
    _paragrafo(cel_data, "DATA", negrito=True, alinhamento=WD_ALIGN_PARAGRAPH.CENTER)
    data_memo = (dados_memo.get("data") or "").strip()
    if data_memo:
        _paragrafo(cel_data, data_memo, alinhamento=WD_ALIGN_PARAGRAPH.CENTER)

    _paragrafo(cel_assinatura, "ASSINATURA", negrito=True, alinhamento=WD_ALIGN_PARAGRAPH.CENTER)
    if dados_escola.get("assinatura_path"):
        _imagem_centralizada(cel_assinatura, dados_escola["assinatura_path"], 3.0)
    else:
        # Sem imagem cadastrada: usa o nome digitado como alternativa.
        assinado_por = (dados_memo.get("assinado_por") or "").strip()
        if assinado_por:
            cel_assinatura.add_paragraph()
            _paragrafo(cel_assinatura, assinado_por, negrito=True, alinhamento=WD_ALIGN_PARAGRAPH.CENTER)

    _paragrafo(cel_recebido, "RECEBIDO POR", negrito=True, alinhamento=WD_ALIGN_PARAGRAPH.CENTER)
    _paragrafo(cel_data2, "DATA", negrito=True, alinhamento=WD_ALIGN_PARAGRAPH.CENTER)

    for c in (cel_data, cel_assinatura, cel_recebido, cel_data2):
        c.add_paragraph()
        c.add_paragraph()
        c.add_paragraph()

    document.save(caminho_saida)
    return caminho_saida


def _config_secao_oficio(document):
    secao = document.sections[0]
    secao.page_width = Cm(21)
    secao.page_height = Cm(29.7)
    secao.left_margin = Cm(3.0)
    secao.right_margin = Cm(2.0)
    secao.top_margin = Cm(2.0)
    secao.bottom_margin = Cm(2.0)


ALTURA_UTIL_PAGINA_OFICIO_CM = 29.7 - 2.0 - 2.0  # altura da pagina menos margens topo/rodape


def gerar_oficio(dados_escola, dados_oficio, caminho_saida):
    """Gera um Ofício (.docx) no "padrão ofício" - uma carta corrida, sem
    quadros/caixas (diferente do Memorando, que usa uma tabela com bordas).
    Cabeçalho (logo + nome da escola), título, local e data por extenso,
    destinatário, assunto, corpo e fecho/assinatura, igual ao modelo oficial
    usado pela administração pública.

    dados_escola: dict com nome_escola, secretaria, logo_path,
        assinatura_path.
    dados_oficio: dict com numero, ano, protocolo (texto livre, opcional),
        data (ja formatada por extenso, ex: "14 de setembro de 2026"), para
        (nome do destinatario), cargo_destinatario (opcional), assunto,
        saudacao, corpo (texto com quebras de linha), fecho, assinado_por
        (nome de quem assina - so aparece se a escola nao tem imagem de
        assinatura cadastrada), cargo_assinado_por (opcional).
    """
    document = docx.Document()
    _config_secao_oficio(document)
    _set_fonte_padrao(document)
    linhas_usadas = 0

    # --- Cabecalho: logo + nome da escola, centralizados tipo timbre ---
    if dados_escola.get("logo_path"):
        _imagem_centralizada(document, dados_escola["logo_path"], 2.5)
        linhas_usadas += 5  # estimativa da altura da imagem em "linhas"

    for i, linha in enumerate(_linhas_cabecalho(dados_escola)):
        _paragrafo(document, linha, negrito=(i == 0), tamanho=13 if i == 0 else TAM_NORMAL,
                   alinhamento=WD_ALIGN_PARAGRAPH.CENTER)
        linhas_usadas += 1
    linha_contato = _linha_contato(dados_escola)
    if linha_contato:
        _paragrafo(document, linha_contato, tamanho=TAM_NORMAL - 1, alinhamento=WD_ALIGN_PARAGRAPH.CENTER)
        linhas_usadas += 1
    secretaria = (dados_escola.get("secretaria") or "").strip()

    document.add_paragraph()
    linhas_usadas += 1

    # --- Titulo + protocolo ---
    numero = (dados_oficio.get("numero") or "").strip()
    ano = (dados_oficio.get("ano") or "").strip()
    titulo = f"OFÍCIO Nº {numero}/{ano}" if numero else "OFÍCIO Nº"
    if secretaria:
        titulo += f" - {secretaria}"
    _paragrafo(document, titulo, negrito=True, tamanho=TAM_TITULO)
    linhas_usadas += 1

    protocolo = (dados_oficio.get("protocolo") or "").strip()
    if protocolo:
        _paragrafo(document, protocolo, negrito=True)
        linhas_usadas += 1

    document.add_paragraph()
    linhas_usadas += 1

    # --- Destinatario ---
    para = (dados_oficio.get("para") or "").strip()
    if para:
        _paragrafo(document, f"A Sua Senhoria o(a) Senhor(a)")
        _paragrafo(document, para, negrito=True)
        linhas_usadas += 2
        cargo_dest = (dados_oficio.get("cargo_destinatario") or "").strip()
        if cargo_dest:
            _paragrafo(document, cargo_dest)
            linhas_usadas += 1

    document.add_paragraph()
    linhas_usadas += 1

    # --- Assunto ---
    p_assunto = document.add_paragraph()
    r1 = p_assunto.add_run("Assunto: ")
    r1.bold = True
    r1.font.name = FONTE
    r1.font.size = Pt(TAM_NORMAL)
    assunto = (dados_oficio.get("assunto") or "").strip()
    if assunto:
        r2 = p_assunto.add_run(assunto)
        r2.font.name = FONTE
        r2.font.size = Pt(TAM_NORMAL)
    linhas_usadas += _estimar_linhas("Assunto: " + assunto)

    document.add_paragraph()
    document.add_paragraph()
    linhas_usadas += 2

    # --- Corpo ---
    p_saud = document.add_paragraph()
    p_saud.paragraph_format.first_line_indent = Cm(1.25)
    run_saud = p_saud.add_run(dados_oficio.get("saudacao") or "Prezado(a) Senhor(a),")
    run_saud.font.name = FONTE
    run_saud.font.size = Pt(TAM_NORMAL)
    linhas_usadas += 1

    document.add_paragraph()
    linhas_usadas += 1

    corpo = dados_oficio.get("corpo") or ""
    for linha in corpo.split("\n"):
        linha = linha.strip()
        if not linha:
            continue
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(1.25)
        run = p.add_run(linha)
        run.font.name = FONTE
        run.font.size = Pt(TAM_NORMAL)
        linhas_usadas += _estimar_linhas(linha)

    # Empurra o fecho/assinatura pra mais perto do fim da pagina, preenchendo
    # com linhas em branco o espaco que sobrar - textos curtos empurram
    # bastante, textos longos empurram pouco ou nada.
    linhas_reservadas_assinatura = 15 if dados_escola.get("assinatura_path") else 11
    linhas_disponiveis = ALTURA_UTIL_PAGINA_OFICIO_CM / ALTURA_LINHA_CM
    linhas_preenchimento = int(linhas_disponiveis - linhas_usadas - linhas_reservadas_assinatura)
    linhas_preenchimento = max(2, min(linhas_preenchimento, 30))
    for _ in range(linhas_preenchimento):
        document.add_paragraph()

    # --- Fecho e assinatura ---
    # Todos os paragrafos daqui pra frente (fecho, assinatura, data) ficam
    # marcados "manter com o proximo", pra sempre ficarem juntos - ou tudo
    # cabe na pagina atual, ou tudo pula junto pra proxima (nunca fica so a
    # data sozinha numa pagina em branco).
    paragrafos_bloco_final = []

    p_fecho = document.add_paragraph()
    p_fecho.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_fecho = p_fecho.add_run(dados_oficio.get("fecho") or "Atenciosamente,")
    run_fecho.font.name = FONTE
    run_fecho.font.size = Pt(TAM_NORMAL)
    paragrafos_bloco_final.append(p_fecho)

    paragrafos_bloco_final.append(document.add_paragraph())

    if dados_escola.get("assinatura_path"):
        paragrafos_bloco_final.append(_imagem_centralizada(document, dados_escola["assinatura_path"], 3.0))
    else:
        assinado_por = (dados_oficio.get("assinado_por") or "").strip()
        if assinado_por:
            paragrafos_bloco_final.append(
                _paragrafo(document, assinado_por, negrito=True, alinhamento=WD_ALIGN_PARAGRAPH.CENTER))
            cargo_assina = (dados_oficio.get("cargo_assinado_por") or "").strip()
            if cargo_assina:
                paragrafos_bloco_final.append(
                    _paragrafo(document, cargo_assina, alinhamento=WD_ALIGN_PARAGRAPH.CENTER))

    paragrafos_bloco_final.append(document.add_paragraph())
    paragrafos_bloco_final.append(document.add_paragraph())

    # --- Local e data, no final de tudo, alinhado a direita ---
    data = (dados_oficio.get("data") or "").strip()
    if data:
        paragrafos_bloco_final.append(_paragrafo(document, f"{data}.", alinhamento=WD_ALIGN_PARAGRAPH.RIGHT))

    for p in paragrafos_bloco_final[:-1]:
        if p is not None:
            p.paragraph_format.keep_with_next = True

    document.save(caminho_saida)
    return caminho_saida


def _caixa(marcado):
    """Retorna a caixinha de opcao marcada ou vazia, ex: (X) / ( )."""
    return "(X)" if marcado else "( )"


def _linha_com_borda_inferior(destino):
    """Cria uma 'linha' (pra assinatura/preenchimento a mao) usando uma
    borda inferior no paragrafo - mais confiavel que espacos sublinhados,
    que podem ser cortados na conversao pra PDF."""
    if destino.paragraphs and not destino.paragraphs[0].runs:
        p = destino.paragraphs[0]
    else:
        p = destino.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    borda = OxmlElement("w:bottom")
    borda.set(qn("w:val"), "single")
    borda.set(qn("w:sz"), "6")
    borda.set(qn("w:space"), "1")
    borda.set(qn("w:color"), "000000")
    pBdr.append(borda)
    pPr.append(pBdr)
    return p


def gerar_declaracao(dados_escola, dados_decl, caminho_saida):
    """Gera uma Declaração Escolar (.docx), no formato de formulário usado
    pela SEMED Manaus (tabela com os dados da escola + texto com opções
    de caixinha marcadas automaticamente conforme o que foi escolhido no
    app, ao contrário de deixar tudo em branco pra marcar a mao).

    dados_escola: dict com nome_escola, secretaria, logo_path, endereco,
        telefone, email.
    dados_decl: dict com aluno, codigo_sigeam, codigo_tipo ("SIGEAM"/
        "Matrícula"/rotulo customizado), situacao_matricula
        ("esta" ou "foi"), ano_letivo, situacao_curso ("cursa" ou "cursou"),
        serie (texto livre), turma, turno ("matutino"/"vespertino"/
        "noturno"/"intermediario"),
        finalidade ("trabalho"/"transferencia"/"sinetram"/"bolsa_familia"/
        "outros"), finalidade_frequencia (numero, so p/ bolsa_familia),
        finalidade_outros (texto, so p/ outros), status_aluno
        ("promovido"/"retido"/"desistente"/"progressao_parcial"/"cursando"),
        status_desistente_data (so p/ desistente), obs (opcional), data
        (ja formatada por extenso).
    """
    document = docx.Document()
    _config_secao_oficio(document)
    _set_fonte_padrao(document)

    # --- Cabecalho: logo + nome da escola/secretaria, lado a lado ---
    logo_path = dados_escola.get("logo_path")
    linhas_cab = _linhas_cabecalho(dados_escola) or ["Secretaria Municipal de Educação"]
    if logo_path:
        cab = document.add_table(rows=1, cols=2)
        _remover_bordas_tabela(cab)
        _definir_largura_colunas(cab, [2.5, 13.5])
        cel_logo, cel_sec = cab.rows[0].cells
        _imagem_centralizada(cel_logo, logo_path, 2.0)
        cel_sec.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        for i, linha in enumerate(linhas_cab):
            _paragrafo(cel_sec, linha.upper(), negrito=(i == 0), tamanho=13 if i == 0 else TAM_NORMAL)
    else:
        for i, linha in enumerate(linhas_cab):
            _paragrafo(document, linha.upper(), negrito=(i == 0), tamanho=13 if i == 0 else TAM_NORMAL)

    document.add_paragraph()

    # --- Tabela com os dados da escola ---
    linhas_info = [
        ("CMEI/Escola Municipal", dados_escola.get("nome_escola", "")),
        ("Endereço", dados_escola.get("endereco", "")),
        ("Telefone", dados_escola.get("telefone", "")),
        ("Email", dados_escola.get("email", "")),
    ]
    tabela_info = document.add_table(rows=len(linhas_info), cols=1)
    tabela_info.style = "Table Grid"
    for linha_tabela, (rotulo, valor) in zip(tabela_info.rows, linhas_info):
        celula = _celula_vazia(linha_tabela.cells[0])
        celula.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = celula.paragraphs[0]
        r1 = p.add_run(f"{rotulo}: ")
        r1.bold = True
        r1.font.name = FONTE
        r1.font.size = Pt(TAM_NORMAL)
        r2 = p.add_run(valor)
        r2.font.name = FONTE
        r2.font.size = Pt(TAM_NORMAL)

    document.add_paragraph()

    # --- Titulo ---
    _paragrafo(document, "DECLARAÇÃO", negrito=True, tamanho=TAM_TITULO, alinhamento=WD_ALIGN_PARAGRAPH.CENTER)

    document.add_paragraph()

    # --- Corpo: texto corrido com as opcoes marcadas ---
    aluno = (dados_decl.get("aluno") or "").strip()
    codigo_sigeam = (dados_decl.get("codigo_sigeam") or "").strip()
    codigo_tipo = (dados_decl.get("codigo_tipo") or "SIGEAM").strip()
    ano_letivo = (dados_decl.get("ano_letivo") or "").strip()
    turma = (dados_decl.get("turma") or "").strip()
    situacao_matricula = dados_decl.get("situacao_matricula") or "esta"
    situacao_curso = dados_decl.get("situacao_curso") or "cursa"
    serie = (dados_decl.get("serie") or "").strip()
    turno = dados_decl.get("turno") or ""

    p_corpo = document.add_paragraph()
    p_corpo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    texto_corpo = (
        f"Declaramos para os devidos fins que o(a) aluno(a), {aluno or '_' * 45}, "
        f"sob o código ({codigo_tipo}) {codigo_sigeam or '_' * 15}, "
        f"{_caixa(situacao_matricula == 'esta')} está {_caixa(situacao_matricula == 'foi')} foi "
        f"matriculado(a) neste estabelecimento de ensino no ano letivo de {ano_letivo or '20____'}, onde "
        f"{_caixa(situacao_curso == 'cursa')} cursa {_caixa(situacao_curso == 'cursou')} cursou a série "
        f"{serie or '_' * 20}, na turma "
        f"{turma or '______'}, turno {_caixa(turno == 'matutino')} matutino {_caixa(turno == 'vespertino')} "
        f"vespertino {_caixa(turno == 'noturno')} noturno {_caixa(turno == 'intermediario')} intermediário, "
        f"conforme especificações abaixo:"
    )
    run_corpo = p_corpo.add_run(texto_corpo)
    run_corpo.font.name = FONTE
    run_corpo.font.size = Pt(TAM_NORMAL)

    document.add_paragraph()

    # --- Duas colunas: finalidade da declaracao + status do aluno ---
    finalidade = dados_decl.get("finalidade") or ""
    freq = (dados_decl.get("finalidade_frequencia") or "").strip()
    outros_texto = (dados_decl.get("finalidade_outros") or "").strip()
    status = dados_decl.get("status_aluno") or ""
    status_data = (dados_decl.get("status_desistente_data") or "").strip()

    tabela_opcoes = document.add_table(rows=6, cols=2)
    _remover_bordas_tabela(tabela_opcoes)
    _definir_largura_colunas(tabela_opcoes, [8.0, 8.0])

    _paragrafo(tabela_opcoes.cell(0, 0), "Declarações para fins de:", negrito=True)
    _paragrafo(tabela_opcoes.cell(1, 0), f"{_caixa(finalidade == 'trabalho')} Trabalho")
    _paragrafo(tabela_opcoes.cell(2, 0), f"{_caixa(finalidade == 'transferencia')} Transferência")
    _paragrafo(tabela_opcoes.cell(3, 0), f"{_caixa(finalidade == 'sinetram')} Sinetram")
    _paragrafo(tabela_opcoes.cell(4, 0),
               f"{_caixa(finalidade == 'bolsa_familia')} Bolsa Família / Frequência {freq or '_____'}")
    _paragrafo(tabela_opcoes.cell(5, 0), f"{_caixa(finalidade == 'outros')} Outros {outros_texto or '_' * 20}")

    _paragrafo(tabela_opcoes.cell(0, 1), "Status do Aluno:", negrito=True)
    _paragrafo(tabela_opcoes.cell(1, 1), f"{_caixa(status == 'promovido')} Promovido (a)")
    _paragrafo(tabela_opcoes.cell(2, 1), f"{_caixa(status == 'retido')} Retido (a)")
    _paragrafo(tabela_opcoes.cell(3, 1),
               f"{_caixa(status == 'desistente')} Desistente a partir de {status_data or '__/__/__'}")
    _paragrafo(tabela_opcoes.cell(4, 1), f"{_caixa(status == 'progressao_parcial')} Progressão Parcial")
    _paragrafo(tabela_opcoes.cell(5, 1), f"{_caixa(status == 'cursando')} Cursando")

    # --- OBS ---
    obs = (dados_decl.get("obs") or "").strip()
    p_obs = document.add_paragraph()
    r1 = p_obs.add_run("OBS: ")
    r1.bold = True
    r1.font.name = FONTE
    r1.font.size = Pt(TAM_NORMAL)
    if obs:
        r2 = p_obs.add_run(obs)
        r2.font.name = FONTE
        r2.font.size = Pt(TAM_NORMAL)
    else:
        _linha_com_borda_inferior(document)

    document.add_paragraph()

    # --- Nota legal fixa ---
    p_nota = document.add_paragraph()
    run_nota = p_nota.add_run(
        "Conforme o parágrafo único do art. 141 do Regimento Geral das Escolas da Rede Municipal, "
        "esta declaração tem validade de 30 dias a contar da data de sua expedição.")
    run_nota.font.name = FONTE
    run_nota.font.size = Pt(TAM_NORMAL - 1)
    run_nota.italic = True

    document.add_paragraph()

    # --- Data ---
    data = (dados_decl.get("data") or "").strip()
    if data:
        p_data = document.add_paragraph()
        run_data = p_data.add_run(f"{data}.")
        run_data.font.name = FONTE
        run_data.font.size = Pt(TAM_NORMAL)
        run_data.bold = True
        run_data.italic = True

    document.add_paragraph()
    document.add_paragraph()

    # --- Assinaturas: Secretario(a) e Diretor(a), lado a lado ---
    tabela_assinatura = document.add_table(rows=2, cols=2)
    _remover_bordas_tabela(tabela_assinatura)
    _definir_largura_colunas(tabela_assinatura, [8.0, 8.0])

    _linha_com_borda_inferior(tabela_assinatura.cell(0, 0))
    _linha_com_borda_inferior(tabela_assinatura.cell(0, 1))

    p_sec = _paragrafo(tabela_assinatura.cell(1, 0), "Secretário (a)", alinhamento=WD_ALIGN_PARAGRAPH.CENTER)
    p_sec.runs[0].italic = True
    p_dir = _paragrafo(tabela_assinatura.cell(1, 1), "Diretor (a)", alinhamento=WD_ALIGN_PARAGRAPH.CENTER)
    p_dir.runs[0].italic = True

    document.save(caminho_saida)
    return caminho_saida
